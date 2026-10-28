"""Word (.docx) reader based on python-docx and the raw WordprocessingML.

* Body content is read in document order straight from the body XML, so
  paragraphs and tables stay interleaved; content controls (``w:sdt``) and
  custom XML wrappers are descended into.
* Run text: ``w:t``, tabs/line breaks as spaces, non-breaking hyphens;
  field codes (``w:instrText``) and tracked deletions are skipped, tracked
  insertions are kept (the text as it looks with all changes accepted).
  Drawings/objects are not part of the paragraph text; text boxes become
  separate paragraph blocks (read by default). Nothing embedded is ever executed.
* Headings: paragraph style (or a style it is based on) named Heading 1–9,
  «Overskrift 1–9» (also as style id ``Overskrift1``) or Title/«Tittel»;
  otherwise the outline level of the paragraph/style. When the document uses
  a Title style, Title becomes level 1 and Heading n level n+1.
* Lists: paragraphs with numbering (``w:numPr``, also via style) become
  ``list_item``; decimal/letter/roman numbers are rendered as a prefix
  («1.», «a)»), bullets are not.
* Tables: one ``table`` block per table; the first row is the header when it
  is marked as a repeating header row or all its text is bold. Merged/nested/
  ragged tables are marked unreliable. A 1×1 table (a layout box) is read as
  normal paragraphs.
* Footnotes/endnotes (``word/footnotes.xml``/``endnotes.xml``, parsed with
  the hardened parser) become ``footnote`` blocks at the end of the
  document: footnotes first, then endnotes, each in reference order.
  The reference point in the body gets a
  visible marker (footnotes 1, 2, 3 …; endnotes i, ii, iii …) covered by a
  ``footnote_ref`` span.
"""

from __future__ import annotations

import re
import threading
from collections.abc import Iterator
from pathlib import Path

from lxml import etree

from ..errors import DocumentImportError
from .cleanup import is_caption, normalize_text, normalize_with_spans
from .common import MSG_EMPTY_DOCUMENT, Progress, ReadResult, check_cancel
from .model import Block, Span
from .postprocess import build_table_data, table_text
from .xmlsafe import UnsafeXMLError, localname, parse_xml

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
EP_NS = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
RT_FOOTNOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
RT_ENDNOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes"


def _w(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


W_P, W_TBL, W_TR, W_TC, W_R, W_T = _w("p"), _w("tbl"), _w("tr"), _w("tc"), _w("r"), _w("t")
W_SDT, W_SDTCONTENT, W_CUSTOMXML = _w("sdt"), _w("sdtContent"), _w("customXml")
W_PPR, W_RPR, W_PSTYLE, W_RSTYLE = _w("pPr"), _w("rPr"), _w("pStyle"), _w("rStyle")
W_VAL, W_ID = _w("val"), _w("id")
W_TXBX = _w("txbxContent")

MSG_DAMAGED = "Word-dokumentet ser ut til å være skadet og kan ikke åpnes."
WARN_TRACKED = (
    "Dokumentet inneholder sporede endringer. Medleser bruker teksten slik den blir "
    "når alle endringene er godtatt."
)
WARN_TABLE = "Noen tabeller har en uklar struktur. Kontroller dem før de leses opp."

_HEADING_RE = re.compile(
    r"^(?:heading|overskrift|rubrik|überschrift|titre|título|kop|otsikko)\s*([1-9])$", re.IGNORECASE
)
_TITLE_NAMES = {"title", "tittel", "titel", "titre", "título", "otsikko"}
_QUOTE_NAMES = {
    "quote", "intense quote", "intensequote", "sitat", "sterkt sitat", "sterktsitat",
    "block text", "blocktext", "blokktekst", "citat", "zitat",
}
_CAPTION_NAMES = {"caption", "bildetekst", "billedtekst", "bildtext", "beskrivelse", "beschriftung", "légende"}

# Containers whose children are walked for paragraph text.
_TEXT_CONTAINERS = {
    _w("hyperlink"), _w("smartTag"), _w("ins"), _w("moveTo"), _w("fldSimple"), W_CUSTOMXML,
    W_SDT, W_SDTCONTENT, _w("dir"), _w("bdo"),
}
_TRACKED_SKIP = {_w("del"), _w("moveFrom")}


def _on(el: etree._Element | None) -> bool:
    """OOXML boolean property: present and not switched off."""
    if el is None:
        return False
    return el.get(W_VAL, "true").lower() not in ("0", "false", "off", "none")


def _roman(n: int) -> str:
    vals = ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
            (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"))
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


def _letters(n: int) -> str:
    # Word style: a..z, aa..zz, …
    n = max(1, n)
    return chr(ord("a") + (n - 1) % 26) * ((n - 1) // 26 + 1)


def _format_number(n: int, fmt: str) -> str | None:
    fmt = (fmt or "decimal").lower()
    if fmt in ("bullet", "none"):
        return None
    if fmt == "lowerletter":
        return _letters(n)
    if fmt == "upperletter":
        return _letters(n).upper()
    if fmt == "lowerroman":
        return _roman(n)
    if fmt == "upperroman":
        return _roman(n).upper()
    if fmt == "decimalzero":
        return f"{n:02d}"
    return str(n)


class _Styles:
    """Paragraph/character style lookup with ``basedOn`` inheritance."""

    def __init__(self, root: etree._Element | None) -> None:
        self.info: dict[str, dict] = {}
        self.default_para: str | None = None
        if root is None:
            return
        for st in root.iter(_w("style")):
            sid = st.get(_w("styleId")) or ""
            name_el = st.find(_w("name"))
            name = name_el.get(W_VAL) if name_el is not None else sid
            based = st.find(_w("basedOn"))
            ppr = st.find(W_PPR)
            rpr = st.find(W_RPR)
            outline = None
            numpr = None
            if ppr is not None:
                ol = ppr.find(_w("outlineLvl"))
                if ol is not None and (ol.get(W_VAL) or "").isdigit():
                    outline = int(ol.get(W_VAL))
                np_el = ppr.find(_w("numPr"))
                if np_el is not None:
                    numpr = np_el
            bold = None
            if rpr is not None and rpr.find(_w("b")) is not None:
                bold = _on(rpr.find(_w("b")))
            self.info[sid] = {
                "name": name or sid,
                "type": st.get(_w("type")),
                "based": based.get(W_VAL) if based is not None else None,
                "outline": outline,
                "numpr": numpr,
                "bold": bold,
            }
            if st.get(_w("type")) == "paragraph" and (st.get(_w("default")) or "").lower() in ("1", "true", "on"):
                self.default_para = sid

    def chain(self, sid: str | None) -> Iterator[tuple[str, dict]]:
        seen: set[str] = set()
        while sid and sid not in seen and len(seen) < 25:
            seen.add(sid)
            info = self.info.get(sid)
            if info is None:
                yield sid, {"name": sid}
                return
            yield sid, info
            sid = info.get("based")

    @staticmethod
    def _keys(sid: str, info: dict) -> list[str]:
        return [str(info.get("name") or "").strip().lower(), sid.strip().lower()]

    def heading_level(self, sid: str | None) -> int | None:
        """0 for Title, 1..9 for Heading n, None otherwise."""
        for s, info in self.chain(sid):
            for key in self._keys(s, info):
                m = _HEADING_RE.match(key)
                if m:
                    return int(m.group(1))
                if key in _TITLE_NAMES:
                    return 0
        return None

    def outline_level(self, sid: str | None) -> int | None:
        for _, info in self.chain(sid):
            if info.get("outline") is not None:
                return info["outline"]
        return None

    def has_name(self, sid: str | None, names: set[str]) -> bool:
        for s, info in self.chain(sid):
            keys = self._keys(s, info)
            if any(k in names or k.replace(" ", "") in names for k in keys):
                return True
        return False

    def numpr(self, sid: str | None) -> etree._Element | None:
        for _, info in self.chain(sid):
            if info.get("numpr") is not None:
                return info["numpr"]
        return None

    def bold(self, sid: str | None) -> bool:
        for _, info in self.chain(sid):
            if info.get("bold") is not None:
                return bool(info["bold"])
        return False


class _Numbering:
    """Minimal list-number renderer from numbering.xml."""

    def __init__(self, root: etree._Element | None) -> None:
        self.abstract: dict[str, dict[int, dict]] = {}
        self.nums: dict[str, tuple[str, dict[int, int]]] = {}
        self.counters: dict[str, dict[int, int]] = {}
        if root is None:
            return
        for an in root.iter(_w("abstractNum")):
            aid = an.get(_w("abstractNumId"))
            levels: dict[int, dict] = {}
            for lvl in an.iter(_w("lvl")):
                try:
                    ilvl = int(lvl.get(_w("ilvl")) or 0)
                except ValueError:
                    continue
                start = lvl.find(_w("start"))
                fmt = lvl.find(_w("numFmt"))
                text = lvl.find(_w("lvlText"))
                levels[ilvl] = {
                    "start": int(start.get(W_VAL)) if start is not None and (start.get(W_VAL) or "").isdigit() else 1,
                    "fmt": fmt.get(W_VAL) if fmt is not None else "decimal",
                    "text": text.get(W_VAL) if text is not None else f"%{ilvl + 1}.",
                }
            if aid is not None:
                self.abstract[aid] = levels
        for num in root.iter(_w("num")):
            nid = num.get(_w("numId"))
            an = num.find(_w("abstractNumId"))
            overrides: dict[int, int] = {}
            for ov in num.iter(_w("lvlOverride")):
                so = ov.find(_w("startOverride"))
                if so is not None and (so.get(W_VAL) or "").isdigit():
                    try:
                        overrides[int(ov.get(_w("ilvl")) or 0)] = int(so.get(W_VAL))
                    except ValueError:
                        pass
            if nid is not None and an is not None:
                self.nums[nid] = (an.get(W_VAL) or "", overrides)

    def prefix(self, num_id: str, ilvl: int) -> str | None:
        """Advance the counter for (num_id, ilvl) and return the rendered number."""
        entry = self.nums.get(num_id)
        if entry is None:
            return None
        aid, overrides = entry
        levels = self.abstract.get(aid, {})
        lvl = levels.get(ilvl)
        if lvl is None:
            return None
        counters = self.counters.setdefault(num_id, {})
        start = overrides.get(ilvl, lvl["start"])
        counters[ilvl] = counters.get(ilvl, start - 1) + 1
        for deeper in [k for k in counters if k > ilvl]:
            del counters[deeper]
        if _format_number(counters[ilvl], lvl["fmt"]) is None:
            return None

        def sub(m: re.Match[str]) -> str:
            k = int(m.group(1)) - 1
            lv = levels.get(k, {"start": 1, "fmt": "decimal"})
            value = counters.get(k, overrides.get(k, lv["start"]))
            return _format_number(value, lv["fmt"]) or ""

        text = re.sub(r"%([1-9])", sub, lvl["text"] or "").strip()
        return text or None


class _Notes:
    """Footnotes/endnotes: id → element, numbered in order of first reference."""

    def __init__(self, doc) -> None:  # noqa: ANN001 - python-docx Document
        self.parts: dict[str, dict[str, etree._Element]] = {"footnote": {}, "endnote": {}}
        self.order: list[tuple[str, str, str]] = []
        self.assigned: dict[tuple[str, str], str] = {}
        self.counters = {"footnote": 0, "endnote": 0}
        for rel in doc.part.rels.values():
            if getattr(rel, "is_external", False):
                continue
            kind = {RT_FOOTNOTES: "footnote", RT_ENDNOTES: "endnote"}.get(rel.reltype)
            if not kind:
                continue
            try:
                root = parse_xml(rel.target_part.blob)
            except (UnsafeXMLError, AttributeError, ValueError):
                continue
            table = {}
            for note in root:
                if localname(note.tag) not in ("footnote", "endnote"):
                    continue
                if note.get(_w("type")) in ("separator", "continuationSeparator", "continuationNotice"):
                    continue
                nid = note.get(W_ID)
                if nid is not None:
                    table[nid] = note
            self.parts[kind] = table

    def ref(self, kind: str, nid: str | None) -> str | None:
        if nid is None or nid not in self.parts[kind]:
            return None
        key = (kind, nid)
        if key in self.assigned:
            return self.assigned[key]
        self.counters[kind] += 1
        n = self.counters[kind]
        marker = str(n) if kind == "footnote" else _roman(n)
        self.assigned[key] = marker
        self.order.append((kind, nid, marker))
        return marker


def _iter_blocks(container: etree._Element) -> Iterator[etree._Element]:
    """Paragraphs and tables of a container in document order."""
    for child in container:
        tag = child.tag
        if tag in (W_P, W_TBL):
            yield child
        elif tag == W_SDT:
            content = child.find(W_SDTCONTENT)
            if content is not None:
                yield from _iter_blocks(content)
        elif tag in (W_CUSTOMXML, _w("ins"), _w("moveTo")):
            yield from _iter_blocks(child)


class _DocxBuilder:
    def __init__(self, doc) -> None:  # noqa: ANN001
        self.doc = doc
        try:
            styles_root = doc.styles.element
        except Exception:  # noqa: BLE001 - missing/odd styles part
            styles_root = None
        self.styles = _Styles(styles_root)
        try:
            numbering_root = doc.part.numbering_part.element
        except Exception:  # noqa: BLE001 - no numbering part
            numbering_root = None
        self.numbering = _Numbering(numbering_root)
        self.notes = _Notes(doc)
        self.blocks: list[Block] = []
        self.tracked = False
        self.unreliable_tables = 0
        self.title_shift = 0

    # -- text ------------------------------------------------------------
    def para_text(self, p: etree._Element) -> tuple[str, list[Span]]:
        parts: list[str] = []
        spans: list[Span] = []
        length = 0

        def add(s: str) -> None:
            nonlocal length
            if s:
                parts.append(s)
                length += len(s)

        def note(kind: str, el: etree._Element) -> None:
            if el.get(_w("customMarkFollows")) in ("1", "true", "on"):
                return
            marker = self.notes.ref(kind, el.get(W_ID))
            if marker:
                start = length
                add(marker)
                spans.append(Span(start, start + len(marker), "footnote_ref"))

        def run(r: etree._Element) -> None:
            for child in r:
                tag = child.tag
                if tag == W_T:
                    add(child.text or "")
                elif tag in (_w("tab"), _w("ptab"), _w("br"), _w("cr")):
                    add(" ")
                elif tag == _w("noBreakHyphen"):
                    add("-")
                elif tag == _w("footnoteReference"):
                    note("footnote", child)
                elif tag == _w("endnoteReference"):
                    note("endnote", child)
                elif tag == _w("delText"):
                    self.tracked = True

        def math(el: etree._Element) -> None:
            for t in el.iter(f"{{{M_NS}}}t"):
                add(t.text or "")
            add(" ")

        def walk(el: etree._Element) -> None:
            for child in el:
                tag = child.tag
                if tag == W_R:
                    run(child)
                elif tag in _TEXT_CONTAINERS:
                    if tag in (_w("ins"), _w("moveTo")):
                        self.tracked = True
                    walk(child)
                elif tag in _TRACKED_SKIP:
                    self.tracked = True
                elif tag in (f"{{{M_NS}}}oMath", f"{{{M_NS}}}oMathPara"):
                    math(child)

        walk(p)
        return "".join(parts), spans

    def plain_para_text(self, p: etree._Element) -> str:
        text, _ = self.para_text(p)
        return normalize_text(text)

    # -- paragraphs ----------------------------------------------------------
    def style_id(self, p: etree._Element) -> str | None:
        ppr = p.find(W_PPR)
        if ppr is not None:
            ps = ppr.find(W_PSTYLE)
            if ps is not None and ps.get(W_VAL):
                return ps.get(W_VAL)
        return self.styles.default_para

    def numbering_of(self, p: etree._Element, sid: str | None) -> tuple[str, int] | None:
        ppr = p.find(W_PPR)
        np_el = ppr.find(_w("numPr")) if ppr is not None else None
        style_np = self.styles.numpr(sid)
        num_id = None
        ilvl = None
        for src in (np_el, style_np):
            if src is None:
                continue
            nid = src.find(_w("numId"))
            lvl = src.find(_w("ilvl"))
            if num_id is None and nid is not None:
                num_id = nid.get(W_VAL)
            if ilvl is None and lvl is not None and (lvl.get(W_VAL) or "").isdigit():
                ilvl = int(lvl.get(W_VAL))
        if num_id is None or num_id == "0":
            return None
        return num_id, ilvl or 0

    def classify(self, p: etree._Element, sid: str | None, text: str) -> tuple[str, int]:
        level = self.styles.heading_level(sid)
        if level is not None:
            return "heading", 1 if level == 0 else min(6, level + self.title_shift)
        ppr = p.find(W_PPR)
        outline = None
        if ppr is not None:
            ol = ppr.find(_w("outlineLvl"))
            if ol is not None and (ol.get(W_VAL) or "").isdigit():
                outline = int(ol.get(W_VAL))
        if outline is None:
            outline = self.styles.outline_level(sid)
        if outline is not None and outline < 9:
            return "heading", min(6, outline + 1 + self.title_shift)
        if self.styles.has_name(sid, _QUOTE_NAMES):
            return "quote", 0
        if self.styles.has_name(sid, _CAPTION_NAMES) or is_caption(text):
            return "caption", 0
        if self.numbering_of(p, sid) is not None:
            return "list_item", 0
        return "paragraph", 0

    def add_paragraph(self, p: etree._Element, *, kind_override: str | None = None) -> None:
        sid = self.style_id(p)
        raw, spans = self.para_text(p)
        text, spans = normalize_with_spans(raw, spans)
        if text:
            kind, level = self.classify(p, sid, text)
            if kind_override:
                kind, level = kind_override, 0
            if kind == "list_item":
                num = self.numbering_of(p, sid)
                prefix = self.numbering.prefix(*num) if num else None
                if prefix:
                    shift = len(prefix) + 1
                    text = f"{prefix} {text}"
                    spans = [Span(s.start + shift, s.end + shift, s.kind) for s in spans]
            self.blocks.append(Block(id="", kind=kind, text=text, level=level, spans=spans))
        if kind_override is None:
            self.add_textboxes(p)

    def add_textboxes(self, p: etree._Element) -> None:
        for box in p.iter(W_TXBX):
            ancestors = [localname(a.tag) for a in box.iterancestors()]
            if "fallback" in ancestors:
                continue  # VML duplicate of the DrawingML text box
            if any(a == "txbxcontent" for a in ancestors):
                continue  # nested text box, handled with its parent
            # Text boxes in Word usually carry real content (call-outs, summaries), so they are
            # read by default like ordinary paragraphs; the user can skip them in the preview.
            for el in _iter_blocks(box):
                if el.tag == W_P:
                    self.add_paragraph(el, kind_override="paragraph")
                elif el.tag == W_TBL:
                    text = self.flatten_table(el)
                    if text:
                        self.blocks.append(Block(id="", kind="paragraph", text=text))

    # -- tables --------------------------------------------------------------
    def cell_text(self, tc: etree._Element) -> tuple[str, bool, bool]:
        """(text, all_bold, has_nested_table)"""
        texts: list[str] = []
        bold_all = True
        nested = False
        for el in _iter_blocks(tc):
            if el.tag == W_P:
                t = self.plain_para_text(el)
                if t:
                    texts.append(t)
                    if not self.para_bold(el):
                        bold_all = False
            elif el.tag == W_TBL:
                nested = True
                t = self.flatten_table(el)
                if t:
                    texts.append(t)
                    bold_all = False
        return " ".join(texts), bold_all, nested

    def para_bold(self, p: etree._Element) -> bool:
        style_bold = self.styles.bold(self.style_id(p))
        saw_text = False
        for r in p.iter(W_R):
            t = "".join(x.text or "" for x in r.iter(W_T))
            if not t.strip():
                continue
            saw_text = True
            rpr = r.find(W_RPR)
            bold = style_bold
            if rpr is not None:
                b = rpr.find(_w("b"))
                if b is not None:
                    bold = _on(b)
                else:
                    rs = rpr.find(W_RSTYLE)
                    if rs is not None and rs.get(W_VAL):
                        rid = rs.get(W_VAL)
                        bold = self.styles.bold(rid) or rid.lower() in ("strong", "sterk", "fet")
            if not bold:
                return False
        return saw_text

    def flatten_table(self, tbl: etree._Element) -> str:
        rows = []
        for tr in tbl.iter(W_TR):
            cells = []
            for tc in tr.iterchildren(W_TC):
                cells.append(" ".join(self.plain_para_text(p) for p in tc.iter(W_P)).strip())
            if any(cells):
                rows.append(" | ".join(cells))
        return normalize_text(" ".join(rows))

    def add_table(self, tbl: etree._Element) -> None:
        trs = [tr for tr in tbl if tr.tag == W_TR]
        for child in tbl:
            if child.tag == W_SDT:
                content = child.find(W_SDTCONTENT)
                if content is not None:
                    trs.extend(tr for tr in content if tr.tag == W_TR)
        rows: list[list[str]] = []
        header_marked: list[bool] = []
        bold_rows: list[bool] = []
        merged = nested = False
        for tr in trs:
            trpr = tr.find(_w("trPr"))
            header_marked.append(trpr is not None and _on(trpr.find(_w("tblHeader"))))
            cells: list[str] = []
            row_bold = True
            row_has_text = False
            tcs = [tc for tc in tr if tc.tag == W_TC]
            for child in tr:
                if child.tag == W_SDT:
                    content = child.find(W_SDTCONTENT)
                    if content is not None:
                        tcs.extend(tc for tc in content if tc.tag == W_TC)
            for tc in tcs:
                tcpr = tc.find(_w("tcPr"))
                span = 1
                if tcpr is not None:
                    gs = tcpr.find(_w("gridSpan"))
                    if gs is not None and (gs.get(W_VAL) or "").isdigit():
                        span = max(1, int(gs.get(W_VAL)))
                    if tcpr.find(_w("vMerge")) is not None or tcpr.find(_w("hMerge")) is not None:
                        merged = True
                if span > 1:
                    merged = True
                text, bold, has_nested = self.cell_text(tc)
                nested = nested or has_nested
                if text:
                    row_has_text = True
                    if not bold:
                        row_bold = False
                cells.append(text)
                cells.extend([""] * (span - 1))
            rows.append(cells)
            bold_rows.append(row_bold and row_has_text)
        if not rows or not any(any(c for c in r) for r in rows):
            return
        if len(rows) == 1 and len(rows[0]) == 1:
            # 1×1 table = layout box: read its content as normal text
            for tc in (tc for tr in trs for tc in tr if tc.tag == W_TC):
                for el in _iter_blocks(tc):
                    if el.tag == W_P:
                        self.add_paragraph(el)
                    elif el.tag == W_TBL:
                        self.add_table(el)
            return
        data = build_table_data(
            rows, first_row_is_header=header_marked[0] or bold_rows[0], merged=merged, nested=nested
        )
        if data is None:
            return
        if not data.reliable:
            self.unreliable_tables += 1
        self.blocks.append(Block(id="", kind="table", text=table_text(data), table=data))

    # -- notes ---------------------------------------------------------------
    def add_notes(self) -> None:
        # footnotes (1, 2, 3 …) first, then endnotes (i, ii, iii …), each in reference order
        ordered = sorted(self.notes.order, key=lambda entry: entry[0] != "footnote")
        for kind, nid, marker in ordered:
            el = self.notes.parts[kind].get(nid)
            if el is None:
                continue
            texts = []
            for p in el.iter(W_P):
                t = self.plain_para_text(p)
                if t:
                    texts.append(t)
            body = " ".join(texts).strip()
            if body:
                self.blocks.append(Block(id="", kind="footnote", text=f"{marker} {body}"))


def _metadata(doc) -> tuple[str | None, dict[str, str]]:  # noqa: ANN001
    meta: dict[str, str] = {}
    title = None
    try:
        cp = doc.core_properties
    except Exception:  # noqa: BLE001
        return None, meta
    fields = {
        "author": "author",
        "subject": "subject",
        "keywords": "keywords",
        "category": "category",
        "last_modified_by": "last_modified_by",
        "language": "language",
        "created": "created",
        "modified": "modified",
    }
    for attr, key in fields.items():
        try:
            value = getattr(cp, attr)
        except Exception:  # noqa: BLE001 - malformed core.xml values
            continue
        if value is None:
            continue
        s = value.isoformat() if hasattr(value, "isoformat") else normalize_text(str(value))
        if s:
            meta[key] = s[:500]
    try:
        title = cp.title or None
    except Exception:  # noqa: BLE001
        title = None
    return title, meta


def _page_count(doc) -> int | None:  # noqa: ANN001
    try:
        for part in doc.part.package.iter_parts():
            if str(part.partname).lower() == "/docprops/app.xml":
                root = parse_xml(part.blob)
                el = root.find(f"{{{EP_NS}}}Pages")
                if el is not None and (el.text or "").strip().isdigit():
                    n = int(el.text.strip())
                    return n if n > 0 else None
    except Exception:  # noqa: BLE001 - optional information only
        return None
    return None


def read_docx(
    path: Path,
    *,
    progress: Progress | None = None,
    cancel: threading.Event | None = None,
) -> ReadResult:
    import docx  # python-docx; imported lazily (heavy)

    report = progress or Progress(None)
    report(0.0, "Åpner Word-dokumentet …")
    try:
        doc = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 - python-docx raises many types
        raise DocumentImportError(MSG_DAMAGED, code="damaged", details=f"python-docx: {exc!r}") from exc

    builder = _DocxBuilder(doc)
    body = doc.element.body
    elements = list(_iter_blocks(body))
    # Title style in use → shift heading levels (Title = 1, Heading 1 = 2 …)
    for el in elements:
        if el.tag == W_P and builder.styles.heading_level(builder.style_id(el)) == 0:
            builder.title_shift = 1
            break

    total = max(1, len(elements))
    for i, el in enumerate(elements):
        if i % 25 == 0:
            check_cancel(cancel)
            report(0.05 + 0.9 * i / total, "Leser Word-dokumentet …")
        if el.tag == W_P:
            builder.add_paragraph(el)
        else:
            builder.add_table(el)
    builder.add_notes()
    check_cancel(cancel)

    if not builder.blocks:
        raise DocumentImportError(MSG_EMPTY_DOCUMENT, code="no_text", details="docx without text")
    title, meta = _metadata(doc)
    result = ReadResult(blocks=builder.blocks, title=title, metadata=meta, page_count=_page_count(doc))
    if builder.tracked:
        result.warn("tracked_changes", WARN_TRACKED)
    if builder.unreliable_tables:
        result.warn("table_uncertain", WARN_TABLE)
    report(1.0, "Word-dokumentet er lest.")
    return result
