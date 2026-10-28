import pytest

from medleser.documents.importer import import_document
from medleser.errors import DocumentImportError

from . import _builders as B


def test_docx_headings_lists_and_table(tmp_path):
    doc = import_document(B.python_docx_document(tmp_path / "oppgave.docx"))
    assert doc.source_type == "docx"
    assert doc.title == "Lesing med lyd"
    assert doc.language == "nb"
    assert doc.metadata["author"] == "Kari Nordmann"
    summary = [(b.kind, b.level, b.text) for b in doc.blocks]
    assert summary[0] == ("heading", 1, "Innledning")
    assert ("list_item", 0, "Første punkt i lista") in summary
    assert ("list_item", 0, "1. Steg en") in summary and ("list_item", 0, "2. Steg to") in summary
    assert ("heading", 2, "Resultater") in summary
    kinds = [b.kind for b in doc.blocks]
    assert kinds.index("table") == kinds.index("heading", 1) + 1  # document order is kept
    table = next(b for b in doc.blocks if b.kind == "table").table
    assert table.header == ["Navn", "Alder", "By"]
    assert table.rows == [["Kari", "21", "Oslo"], ["Ola", "23", "Bergen"]]
    assert table.reliable
    assert doc.blocks[-1].kind == "paragraph"


def test_docx_text_box_content_is_read_by_default(tmp_path):
    """Text boxes often hold real content; they must not be silently skipped."""
    import docx
    from docx.oxml import parse_xml

    from medleser.speech.text_prep import ReadingOptions, build_plan

    d = docx.Document()
    d.add_paragraph("Vanlig avsnitt før boksen.")
    p = d.add_paragraph("Avsnitt med tekstboks.")
    pict = parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape><v:textbox><w:txbxContent>'
        "<w:p><w:r><w:t>Viktig: Husk å levere oppgaven innen fredag.</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
    )
    p._p.append(pict)
    path = tmp_path / "boks.docx"
    d.save(path)
    doc = import_document(path)
    box = next(b for b in doc.blocks if "levere oppgaven" in b.text)
    assert box.kind == "paragraph"
    plan = build_plan(doc, ReadingOptions(), "nb")
    assert any("levere oppgaven" in s.text for s in plan.segments)


def test_docx_norwegian_styles_footnotes_fields_and_tracked_changes(tmp_path):
    doc = import_document(B.handmade_docx(tmp_path / "fotnoter.docx"))
    by_text = {b.text: b for b in doc.blocks}
    assert by_text["Innledning"].kind == "heading" and by_text["Innledning"].level == 1
    assert by_text["Arvet overskrift"].kind == "heading"  # style based on «Overskrift 1»
    assert by_text["Bakgrunn"].level == 2
    para = next(b for b in doc.blocks if b.text.startswith("Dette er et avsnitt"))
    assert para.text == "Dette er et avsnitt med en fotnote1 og mer tekst etterpå."
    ref = [s for s in para.spans if s.kind == "footnote_ref"]
    assert len(ref) == 1 and para.text[ref[0].start:ref[0].end] == "1"
    assert doc.blocks[-1].kind == "footnote"
    assert doc.blocks[-1].text == "1 Dette er fotnoteteksten."
    assert "Lenketekst står her." in by_text  # field result kept, field code skipped
    assert "HYPERLINK" not in doc.plain_text()
    assert "slettet tekst" not in doc.plain_text()
    assert "tracked_changes" in [w.code for w in doc.warnings]
    assert doc.title == "Fotnoter i praksis"


def test_docx_merged_table_is_unreliable(tmp_path):
    body = (
        "<w:tbl>"
        "<w:tr><w:tc><w:p><w:r><w:t>A</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>B</w:t></w:r></w:p></w:tc></w:tr>"
        '<w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr><w:p><w:r><w:t>Sammenslått</w:t></w:r></w:p></w:tc></w:tr>'
        "</w:tbl>"
        "<w:p><w:r><w:t>Avsnitt etter tabellen med litt mer tekst.</w:t></w:r></w:p>"
    )
    doc = import_document(B.handmade_docx(tmp_path / "t.docx", body=body))
    table = next(b for b in doc.blocks if b.kind == "table").table
    assert not table.reliable
    assert "sammenslåtte" in table.note
    assert "table_uncertain" in [w.code for w in doc.warnings]


def test_docx_outline_level_title_style_and_endnotes(tmp_path):
    body = (
        '<w:p><w:pPr><w:pStyle w:val="Title"/></w:pPr><w:r><w:t>Hovedtittel</w:t></w:r></w:p>'
        '<w:p><w:pPr><w:outlineLvl w:val="0"/></w:pPr><w:r><w:t>Kapittel med disposisjonsnivå</w:t></w:r></w:p>'
        '<w:p><w:r><w:t xml:space="preserve">Tekst med sluttnote</w:t></w:r><w:r><w:endnoteReference w:id="2"/></w:r>'
        '<w:r><w:t xml:space="preserve"> og fotnote</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p>'
    )
    endnotes = (
        f'<w:endnotes xmlns:w="{B.W_NS}"><w:endnote w:id="2"><w:p><w:r><w:endnoteRef/></w:r>'
        '<w:r><w:t xml:space="preserve"> Sluttnoten.</w:t></w:r></w:p></w:endnote></w:endnotes>'
    )
    rels = B.DOCX_DOC_RELS.replace(
        "</Relationships>",
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes" '
        'Target="endnotes.xml"/></Relationships>',
    )
    styles = B.DOCX_STYLES.replace(
        "</w:styles>", '<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/></w:style></w:styles>'
    )
    path = B.handmade_docx(tmp_path / "noter.docx", body=body, title="",
                           extra={"word/endnotes.xml": endnotes, "word/_rels/document.xml.rels": rels,
                                  "word/styles.xml": styles})
    doc = import_document(path)
    levels = {b.text: (b.kind, b.level) for b in doc.blocks}
    assert levels["Hovedtittel"] == ("heading", 1)
    assert levels["Kapittel med disposisjonsnivå"] == ("heading", 2)  # shifted below the Title
    para = next(b for b in doc.blocks if b.text.startswith("Tekst med sluttnote"))
    assert para.text == "Tekst med sluttnotei og fotnote1"
    assert [para.text[s.start:s.end] for s in para.spans if s.kind == "footnote_ref"] == ["i", "1"]
    notes = [b.text for b in doc.blocks if b.kind == "footnote"]
    assert notes == ["1 Dette er fotnoteteksten.", "i Sluttnoten."]
    assert doc.title == "Hovedtittel"  # empty core title → first top-level heading


def test_docm_with_macros_is_refused(tmp_path):
    path = B.handmade_docx(
        tmp_path / "makro.docx",
        main_type="application/vnd.ms-word.document.macroEnabled.main+xml",
        extra={"word/vbaProject.bin": b"\xd0\xcf\x11\xe0 fake macro"},
    )
    with pytest.raises(DocumentImportError) as exc:
        import_document(path)
    assert exc.value.code == "docm_macro"
    assert "makro" in exc.value.message
