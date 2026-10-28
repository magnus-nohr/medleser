"""Synthetic test documents, generated on the fly (no binary fixtures in the repo)."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pymupdf

# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

NB_PARAGRAPH = (
    "Studenter med dysleksi bruker ofte lengre tid på å lese faglitteratur. Det er derfor viktig at "
    "de får tilgang til hjelpemidler som kan lese teksten høyt, og at lærerne vet hvordan slike "
    "verktøy fungerer i praksis på høgskolen."
)
NB_PARAGRAPH_2 = (
    "I denne undersøkelsen har vi sett på hvordan studentene opplever lydstøtte, og hva som gjør at "
    "noen velger å bruke den mer enn andre. Svarene viser at både æ, ø og å må uttales riktig."
)


def text_width(text: str, size: float, font: str = "helv") -> float:
    return pymupdf.get_text_length(text, fontname=font, fontsize=size)


def add_lines(page, x: float, y: float, lines: list[str], *, size: float = 11, font: str = "helv",
              leading: float | None = None) -> float:
    """Write lines starting at baseline y; returns the next baseline."""
    leading = leading or size * 1.3
    for line in lines:
        page.insert_text((x, y), line, fontname=font, fontsize=size)
        y += leading
    return y


def wrap(text: str, width: float, size: float = 11, font: str = "helv") -> list[str]:
    words = text.split()
    lines: list[str] = []
    cur = ""
    for w in words:
        cand = f"{cur} {w}".strip()
        if text_width(cand, size, font) <= width:
            cur = cand
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def norwegian_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 90), "Lesing med lyd i høyere utdanning", fontname="hebo", fontsize=20)
    y = add_lines(page, 72, 140, wrap(NB_PARAGRAPH, 450))
    add_lines(page, 72, y + 14, wrap(NB_PARAGRAPH_2, 450))
    doc.set_metadata({"title": "", "author": "Kari Nordmann"})
    doc.save(path)
    return path


LEFT_TEXT = (
    "Alpha studies show that readers with dyslexia benefit from synchronized audio support. "
    "Alpha participants listened while reading highlighted text on the screen during lectures. "
) * 4
RIGHT_TEXT = (
    "Omega results indicate improved comprehension and reduced fatigue among the students. "
    "Omega effects were strongest for long academic texts with complex vocabulary and syntax. "
) * 4


def two_column_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_textbox(pymupdf.Rect(60, 50, 535, 95), "Reading Along: Audio Support in Higher Education",
                        fontsize=18, fontname="hebo", align=1)
    page.insert_textbox(pymupdf.Rect(60, 100, 535, 170),
                        "Abstract. This paper examines synchronized read-along support for students with "
                        "dyslexia across several courses and reports a mixed methods evaluation.",
                        fontsize=10, fontname="helv")
    # The right column is written first in the content stream on purpose.
    page.insert_textbox(pymupdf.Rect(310, 190, 535, 780), RIGHT_TEXT, fontsize=10, fontname="helv")
    page.insert_textbox(pymupdf.Rect(60, 190, 285, 780), LEFT_TEXT, fontsize=10, fontname="helv")
    doc.save(path)
    return path


def header_footer_pdf(path: Path, pages: int = 4) -> Path:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 40), "Tidsskrift for lesing og læring 2024", fontname="helv", fontsize=9)
        body = f"Side {i + 1} handler om lesing. " + NB_PARAGRAPH + " " + NB_PARAGRAPH_2
        add_lines(page, 72, 110, wrap(body, 450))
        page.insert_text((72, 800), "Høgskolen i Østfold", fontname="helv", fontsize=9)
        page.insert_text((470, 800), f"Side {i + 1} av {pages}", fontname="helv", fontsize=9)
    doc.save(path)
    return path


def footnote_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 80), "Bakgrunn", fontname="hebo", fontsize=14)
    first = "Mange studenter bruker lydbøker i studiene"
    page.insert_text((72, 120), first, fontname="helv", fontsize=11)
    w = text_width(first, 11)
    page.insert_text((72 + w + 0.5, 115.5), "1", fontname="helv", fontsize=7)
    page.insert_text((72 + w + 4.5, 120), " og opplever at det hjelper.", fontname="helv", fontsize=11)
    add_lines(page, 72, 134.3, wrap(NB_PARAGRAPH, 450))
    page.insert_text((72, 770), "1 Dette er en fotnote om kildene til undersøkelsen.", fontname="helv", fontsize=8)
    doc.save(path)
    return path


def table_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    y = add_lines(page, 72, 80, wrap("Tabellen nedenfor viser deltakerne i studien og hvor de bor.", 450))
    cols = [72, 200, 330, 460]
    rows = [y + 10 + 20 * k for k in range(4)]
    for x in cols:
        page.draw_line((x, rows[0]), (x, rows[-1]))
    for yy in rows:
        page.draw_line((cols[0], yy), (cols[-1], yy))
    data = [["Navn", "Alder", "By"], ["Kari", "21", "Oslo"], ["Ola", "23", "Bergen"]]
    for r, row in enumerate(data):
        for c, cell in enumerate(row):
            page.insert_text((cols[c] + 4, rows[r] + 14), cell, fontname="hebo" if r == 0 else "helv", fontsize=10)
    add_lines(page, 72, rows[-1] + 30, wrap("Etter tabellen fortsetter teksten som vanlig.", 450))
    doc.save(path)
    return path


REFERENCES = [
    ["Andersen, K. (2019). Lesing og forståelse i høyere utdanning. Universitetsforlaget.",
     "https://doi.org/10.1234/abcd.2019.001"],
    ["Berg, L., & Hansen, T. (2021). Lydbøker som hjelpemiddel. Norsk pedagogisk tidsskrift,",
     "105(2), 45-60. https://www.example.no/artikkel"],
    ["Carlsen, M. (2020). Dysleksi og studier. Hentet fra www.dysleksi.no"],
]


def references_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 80), "Innledning", fontname="hebo", fontsize=14)
    y = add_lines(page, 72, 110, wrap(NB_PARAGRAPH, 450))
    page.insert_text((72, y + 20), "Referanser", fontname="hebo", fontsize=14)
    y += 50
    for entry in REFERENCES:
        for k, line in enumerate(entry):
            page.insert_text((72 if k == 0 else 90, y), line, fontname="helv", fontsize=10)
            y += 13
    doc.save(path)
    return path


HYPHEN_LINES = [
    "Studentene trenger bedre infor-",
    "masjon om hjelpemidler, og IKT-",
    "basert støtte er viktig for lese-",
    "og skrivevansker. Send gjerne en e-",
    "post til veilederen din om dette.",
]


def hyphenation_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    add_lines(page, 72, 100, HYPHEN_LINES)
    doc.save(path)
    return path


def figure_pdf(path: Path) -> Path:
    """An image with labels on top, a caption and a vector bar chart with small labels."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    y = add_lines(page, 72, 80, wrap(NB_PARAGRAPH, 450))
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), 0)
    pix.clear_with(200)
    page.insert_image(pymupdf.Rect(72, y + 10, 372, y + 160), stream=pix.tobytes("png"))
    page.insert_text((90, y + 50), "Lesetid", fontsize=9)
    page.insert_text((250, y + 120), "Forståelse", fontsize=9)
    page.insert_text((72, y + 180), "Figur 1: Sammenheng mellom lesetid og forståelse.", fontsize=9)
    y2 = y + 220
    for k in range(8):
        h = 20 + 10 * k
        page.draw_rect(pymupdf.Rect(80 + k * 30, y2 + 120 - h, 100 + k * 30, y2 + 120), color=(0, 0, 0),
                       fill=(0.5, 0.5, 0.8))
        page.draw_line((80 + k * 30, y2 + 125), (100 + k * 30, y2 + 125))
        page.insert_text((80 + k * 30, y2 + 137), f"G{k}", fontsize=7)
    page.draw_line((75, y2), (75, y2 + 125))
    page.draw_line((75, y2 + 125), (330, y2 + 125))
    add_lines(page, 72, y2 + 170, wrap(NB_PARAGRAPH_2, 450))
    doc.save(path)
    return path


def booktabs_pdf(path: Path) -> Path:
    """A table with only horizontal rules (typical for journal articles)."""
    doc = pymupdf.open()
    p = doc.new_page(width=595, height=842)
    add_lines(p, 72, 70, ["Tabell 2 viser gjennomsnittlige skårer for gruppene i studien."])
    p.insert_text((72, 92), "Tabell 2. Skårer per gruppe", fontsize=10)
    p.draw_line((72, 100), (500, 100), width=1)
    p.draw_line((72, 120), (500, 120), width=0.5)
    p.draw_line((72, 180), (500, 180), width=1)
    rows = [["Gruppe", "Antall", "Snitt"], ["Kontroll", "24", "3.1"], ["Tiltak", "26", "4.2"], ["Total", "50", "3.7"]]
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            p.insert_text((80 + c * 150, [114, 134, 152, 170][r]), cell, fontsize=10, fontname="hebo" if r == 0 else "helv")
    add_lines(p, 72, 210, wrap(NB_PARAGRAPH, 450))
    doc.save(path)
    return path


def merged_table_pdf(path: Path) -> Path:
    """Ruled table without a header row and with a merged last row."""
    doc = pymupdf.open()
    p = doc.new_page(width=595, height=842)
    add_lines(p, 72, 70, ["Innledende tekst om tabellen under."])
    cols = [72, 200]
    rws = [90, 110, 130, 150]
    for x in (72, 330):
        p.draw_line((x, 90), (x, 150))
    p.draw_line((200, 90), (200, 130))
    for yy in rws:
        p.draw_line((72, yy), (330, yy))
    for r, row in enumerate([["Kari", "Oslo"], ["Ola", "Bergen"]]):
        for c, cell in enumerate(row):
            p.insert_text((cols[c] + 4, rws[r] + 14), cell, fontsize=10)
    p.insert_text((76, 144), "Sammenslått rad", fontsize=10)
    doc.save(path)
    return path


STORY_P = (
    "Studenter med dysleksi bruker ofte lengre tid på å lese faglitteratur, og mange opplever at "
    "lange tekster blir krevende å komme gjennom. Lydstøtte med synkron markering kan gjøre det enklere "
    "å holde konsentrasjonen, særlig når teksten inneholder mange fagbegreper og lange setninger. "
)
STORY_PARAGRAPHS = (
    [f"Avsnitt {k}. {STORY_P * 2}".strip() for k in range(1, 4)]
    + [f"Metodeavsnitt {k}. {STORY_P * 3}".strip() for k in range(1, 5)]
    + [f"Diskusjonsavsnitt {k}. {STORY_P * 2}".strip() for k in range(1, 4)]
)


def story_article_pdf(path: Path) -> Path:
    """A realistically typeset two-column article (justified text flowing over columns and pages)."""
    p = STORY_PARAGRAPHS
    html = "<h2>Innledning</h2>" + "".join(f"<p>{t}</p>" for t in p[:3])
    html += "<h2>Metode</h2>" + "".join(f"<p>{t}</p>" for t in p[3:7])
    html += "<h2>Diskusjon</h2>" + "".join(f"<p>{t}</p>" for t in p[7:])
    html += "<h2>Referanser</h2>" + "".join(
        f'<p style="text-indent:-18px;margin-left:18px">Forfatter{k}, A. ({2000 + k}). Tittel på en bok om '
        f"lesing og læring nummer {k}. Universitetsforlaget. https://doi.org/10.1234/bok.{k}</p>"
        for k in range(1, 7)
    )
    css = "p { font-size: 10px; text-align: justify; margin: 0 0 6px 0; } h2 { font-size: 13px; margin: 8px 0 4px 0; }"
    story = pymupdf.Story(html=html, user_css=css)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    columns = [pymupdf.Rect(50, 110, 290, 780), pymupdf.Rect(305, 110, 545, 780)]
    more = True
    while more:
        device = writer.begin_page(pymupdf.Rect(0, 0, 595, 842))
        for col in columns:
            more, _ = story.place(col)
            story.draw(device)
            if not more:
                break
        writer.end_page()
    writer.close()
    doc = pymupdf.open("pdf", buffer.getvalue())
    for i, page in enumerate(doc):
        page.insert_text((50, 40), "Nordisk tidsskrift for lesing 2024", fontsize=8)
        page.insert_text((290, 815), str(i + 1), fontsize=8)
    doc[0].insert_textbox(pymupdf.Rect(50, 55, 545, 100), "Synkron lydstøtte for studenter med dysleksi",
                          fontsize=18, fontname="hebo", align=1)
    doc.save(path)
    return path


def scanned_pdf(path: Path, lines: list[str] | None = None) -> Path:
    """Rasterise a text page into an image-only PDF (no text layer)."""
    if lines is None:
        lines = ["Dette er en skannet side med norsk tekst.", "Tekstgjenkjenning gjør den lesbar."]
    src = pymupdf.open()
    sp = src.new_page(width=595, height=842)
    add_lines(sp, 72, 120, lines, size=18, leading=30)
    pix = sp.get_pixmap(dpi=200)
    out = pymupdf.open()
    page = out.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    out.save(path)
    return path


def simple_pdf(path: Path, text: str = NB_PARAGRAPH, **save_kwargs) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    add_lines(page, 72, 100, wrap(text, 450))
    doc.save(path, **save_kwargs)
    return path


def many_pages_pdf(path: Path, pages: int) -> Path:
    """Two-column pages, copied to build a long document quickly."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((60, 40), "Journal of Testing 2024", fontname="helv", fontsize=8)
    page.insert_textbox(pymupdf.Rect(60, 70, 285, 790), LEFT_TEXT * 2, fontsize=9.5, fontname="helv")
    page.insert_textbox(pymupdf.Rect(310, 70, 535, 790), RIGHT_TEXT * 2, fontsize=9.5, fontname="helv")
    page.insert_text((290, 815), "1", fontname="helv", fontsize=9)
    while doc.page_count < pages:  # doubling: fast to build, real pages to parse
        chunk = pymupdf.open()
        chunk.insert_pdf(doc, from_page=0, to_page=min(doc.page_count, pages - doc.page_count) - 1)
        doc.insert_pdf(chunk)
        chunk.close()
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def python_docx_document(path: Path) -> Path:
    import docx

    d = docx.Document()
    d.core_properties.title = "Lesing med lyd"
    d.core_properties.author = "Kari Nordmann"
    d.add_heading("Innledning", level=1)
    d.add_paragraph(NB_PARAGRAPH)
    d.add_paragraph("Første punkt i lista", style="List Bullet")
    d.add_paragraph("Andre punkt i lista", style="List Bullet")
    d.add_paragraph("Steg en", style="List Number")
    d.add_paragraph("Steg to", style="List Number")
    d.add_heading("Resultater", level=2)
    table = d.add_table(rows=3, cols=3)
    cells = [["Navn", "Alder", "By"], ["Kari", "21", "Oslo"], ["Ola", "23", "Bergen"]]
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            para = table.rows[r].cells[c].paragraphs[0]
            run = para.add_run(value)
            run.bold = r == 0
    d.add_paragraph(NB_PARAGRAPH_2)
    d.save(path)
    return path


def _zip(path: Path, files: dict[str, str | bytes]) -> Path:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    return path


DOCX_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="{main}"/>
  <Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
  <Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>
  <Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
</Types>"""
DOCX_MAIN = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
DOCX_ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
</Relationships>"""
DOCX_DOC_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>
</Relationships>"""
DOCX_CORE = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
  xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/"
  xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <dc:title>{title}</dc:title><dc:creator>Kari Nordmann</dc:creator>
</cp:coreProperties>"""
DOCX_STYLES = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="{W_NS}">
  <w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>
  <w:style w:type="paragraph" w:styleId="Overskrift1"><w:name w:val="Overskrift 1"/><w:basedOn w:val="Normal"/></w:style>
  <w:style w:type="paragraph" w:styleId="Overskrift2"><w:name w:val="Overskrift 2"/><w:basedOn w:val="Normal"/></w:style>
  <w:style w:type="paragraph" w:styleId="MinOverskrift"><w:name w:val="Min overskrift"/><w:basedOn w:val="Overskrift1"/></w:style>
</w:styles>"""
DOCX_FOOTNOTES = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="{W_NS}">
  <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
  <w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
  <w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> Dette er fotnoteteksten.</w:t></w:r></w:p></w:footnote>
</w:footnotes>"""


def handmade_docx(path: Path, *, body: str | None = None, main_type: str = DOCX_MAIN,
                  extra: dict[str, str | bytes] | None = None, title: str = "Fotnoter i praksis") -> Path:
    body = body or (
        '<w:p><w:pPr><w:pStyle w:val="Overskrift1"/></w:pPr><w:r><w:t>Innledning</w:t></w:r></w:p>'
        '<w:p><w:r><w:t xml:space="preserve">Dette er et avsnitt med en fotnote</w:t></w:r>'
        '<w:r><w:footnoteReference w:id="1"/></w:r>'
        '<w:r><w:t xml:space="preserve"> og mer tekst etterpå.</w:t></w:r>'
        '<w:del w:id="9" w:author="x"><w:r><w:delText>slettet tekst</w:delText></w:r></w:del></w:p>'
        '<w:p><w:pPr><w:pStyle w:val="MinOverskrift"/></w:pPr><w:r><w:t>Arvet overskrift</w:t></w:r></w:p>'
        '<w:p><w:pPr><w:pStyle w:val="Overskrift2"/></w:pPr><w:r><w:t>Bakgrunn</w:t></w:r></w:p>'
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>HYPERLINK "https://x.no"</w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>Lenketekst</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t xml:space="preserve"> står her.</w:t></w:r></w:p>'
    )
    document = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="{W_NS}"><w:body>{body}<w:sectPr/></w:body></w:document>'
    files: dict[str, str | bytes] = {
        "[Content_Types].xml": DOCX_CONTENT_TYPES.format(main=main_type),
        "_rels/.rels": DOCX_ROOT_RELS,
        "word/_rels/document.xml.rels": DOCX_DOC_RELS,
        "word/document.xml": document,
        "word/styles.xml": DOCX_STYLES,
        "word/footnotes.xml": DOCX_FOOTNOTES,
        "docProps/core.xml": DOCX_CORE.format(title=title),
    }
    files.update(extra or {})
    return _zip(path, files)


# ---------------------------------------------------------------------------
# EPUB
# ---------------------------------------------------------------------------

CONTAINER = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>"""

CHAPTER_1 = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="nb">
<head><title>Kapittel 1</title><script>var secret = "SKRIPTTEKST";</script><style>p { color: red; }</style></head>
<body>
<h1>Kapittel 1: Begynnelsen</h1>
<p>Studenter med dysleksi bruker ofte lengre tid på å lese&nbsp;faglitteratur.<a epub:type="noteref" href="#n1">1</a> Det er derfor viktig at de får hjelpemidler.</p>
<script>document.write("SKRIPTTEKST")</script>
<ul><li>Første punkt</li><li>Andre punkt</li></ul>
<blockquote><p>Et sitat fra en kjent forsker.</p></blockquote>
<aside epub:type="footnote" id="n1"><p>1 Dette er en fotnote i e-boken.</p></aside>
</body></html>"""

CHAPTER_2 = """<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" lang="nb">
<head><title>Kapittel 2</title></head>
<body>
<h1>Kapittel 2: Fortsettelsen</h1>
<p>I denne undersøkelsen har vi sett på hvordan studentene opplever lydstøtte, og hva som gjør at noen velger å bruke den.</p>
<table><thead><tr><th>Navn</th><th>Alder</th></tr></thead><tbody><tr><td>Kari</td><td>21</td></tr></tbody></table>
<span epub:type="pagebreak" xmlns:epub="http://www.idpf.org/2007/ops" title="12">12</span>
</body></html>"""

NAV = """<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">
<head><title>Innhold</title></head>
<body><nav epub:type="toc"><ol><li><a href="ch1.xhtml">NAVIGASJONSTEKST</a></li></ol></nav></body></html>"""


def make_epub(path: Path, *, title: str = "Lydbok for studenter", chapters: list[tuple[str, str]] | None = None,
              extra: dict[str, str | bytes] | None = None) -> Path:
    chapters = chapters or [("ch1.xhtml", CHAPTER_1), ("ch2.xhtml", CHAPTER_2)]
    items = "\n".join(
        f'<item id="c{i}" href="{name}" media-type="application/xhtml+xml"/>' for i, (name, _) in enumerate(chapters)
    )
    spine = "\n".join(f'<itemref idref="c{i}"/>' for i in range(len(chapters)))
    opf = f"""<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="uid">urn:uuid:12345</dc:identifier>
    <dc:title>{title}</dc:title><dc:creator>Ola Nordmann</dc:creator><dc:language>nb</dc:language>
  </metadata>
  <manifest>
    <item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
    {items}
  </manifest>
  <spine><itemref idref="nav"/>{spine}</spine>
</package>"""
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        zf.writestr("META-INF/container.xml", CONTAINER, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("OEBPS/nav.xhtml", NAV, compress_type=zipfile.ZIP_DEFLATED)
        for name, content in chapters:
            zf.writestr(f"OEBPS/{name}", content, compress_type=zipfile.ZIP_DEFLATED)
        for name, content in (extra or {}).items():
            zf.writestr(name, content, compress_type=zipfile.ZIP_DEFLATED)
    return path


def encryption_xml(algorithm: str, uri: str = "OEBPS/ch1.xhtml") -> str:
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<encryption xmlns="urn:oasis:names:tc:opendocument:xmlns:container"
            xmlns:enc="http://www.w3.org/2001/04/xmlenc#">
  <enc:EncryptedData>
    <enc:EncryptionMethod Algorithm="{algorithm}"/>
    <enc:CipherData><enc:CipherReference URI="{uri}"/></enc:CipherData>
  </enc:EncryptedData>
</encryption>"""
